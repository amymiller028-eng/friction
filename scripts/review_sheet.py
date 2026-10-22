"""Write an 'anger-review' sheet: the Q6 answers whose coding is a judgment call, with a blank column for the
editor's own code. Nothing in the coding changes; this only lists what to review."""
import re
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

XLSX, ROWS = sys.argv[1], sys.argv[2]
r = pd.read_pickle(ROWS)
ans = r['Answer (Q6)'].fillna('').astype(str)
low = ans.str.lower().str.strip()
part = r['Part coded'].fillna('').str.strip()
spec = r['Specific emotions'].fillna('')

# emotions placed by judgment because they are not in the feeling-words table (vulnerability already decided: afraid)
NOT_IN_TABLE = {'Hate', 'Disdain', 'Contempt', 'Resentment', 'Bitterness', 'Weakness', 'Helplessness', 'Overwhelm',
                'Stress', 'Uncertainty', 'Doubt', 'Confusion', 'Love', 'Affection', 'Attraction', 'Humor', 'Hope', 'Shyness', 'Pride'}

reasons = []
for i in r.index:
    why = []
    words = [w for w in spec[i].split(', ') if w]
    if re.match(r'^(crying|cry)\b', low[i]) and r.at[i, 'Named an emotion'] == 'No':
        why.append('Names crying (a behavior). Not counted; could be counted as Sad')
    if r.at[i, 'Why not counted'].startswith('Reviewed by hand'):
        why.append('Set aside by hand as an insult')
    if part[i] and part[i] == low[i] and r.at[i, 'Named an emotion'] == 'Yes' and re.search(r'\b(because|cause|since|when|if)\b|[.!?;:]', low[i]):
        why.append('No emotion word in the part they named, so the whole answer was read')
    odd = [w for w in words if w in NOT_IN_TABLE]
    if odd:
        why.append('Not in the feeling-words table, placed by judgment: ' + ', '.join(odd))
    if r.at[i, 'Why not counted'] == 'No core emotion named' and len(part[i].split()) <= 4:
        why.append('Short answer that may name an emotion outside the five (not counted)')
    if len({c for c in ['Angry', 'Afraid', 'Sad', 'Ashamed', 'Happy'] if r.at[i, c] == 1}) > 1:
        why.append('Counted under more than one core emotion')
    reasons.append('; '.join(why))
r['Why I am unsure'] = reasons
rev = r[r['Why I am unsure'] != ''].copy()
rev['My coding'] = [', '.join(c for c in ['Angry', 'Afraid', 'Sad', 'Ashamed', 'Happy'] if row[c] == 1) or 'Not counted'
                    for _, row in rev.iterrows()]
rev = rev[['ResponseId', 'Answer (Q6)', 'My coding', 'Specific emotions', 'Why I am unsure']]
rev['Your code'] = ''
rev['Notes'] = ''

wb = load_workbook(XLSX)
if 'anger-review' in wb.sheetnames:
    del wb['anger-review']
ws = wb.create_sheet('anger-review')
ws.cell(1, 1, 'Q6 answers to review: which emotion feels most off-limits at work?').font = Font(bold=True, size=14, color='002D61')
ws.cell(2, 1, f'{len(rev)} answers where the coding is a judgment call. Pick a code in "Your code" (blank = keep my coding); '
              'the analysis is re-run from your codes.').font = Font(italic=True, color='5A6A82')
H = 4
for j, name in enumerate(rev.columns, 1):
    c = ws.cell(H, j, name)
    c.fill, c.font = PatternFill('solid', fgColor='002D61'), Font(bold=True, color='FFFFFF')
for i, row in enumerate(rev.itertuples(index=False), H + 1):
    for j, v in enumerate(row, 1):
        ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(i, 6).fill = PatternFill('solid', fgColor='FFF5E4')
dv = DataValidation(type='list', allow_blank=True, showDropDown=False,
                    formula1='"Angry,Afraid,Sad,Ashamed,Happy,Angry + Sad,Angry + Afraid,Afraid + Sad,Not counted"')
dv.add(f'F{H + 1}:F{H + len(rev)}')
ws.add_data_validation(dv)
for col, w in zip('ABCDEFG', [30, 60, 16, 26, 46, 16, 30]):
    ws.column_dimensions[col].width = w
ws.auto_filter.ref = f'A{H}:G{H + len(rev)}'
ws.freeze_panes = f'A{H + 1}'
wb.save(XLSX)
print(len(rev), 'answers to review')
print(rev['Why I am unsure'].str.split('; ').explode().str.split(':').str[0].value_counts().to_string())
