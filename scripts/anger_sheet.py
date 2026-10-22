"""Write the 'anger-at-work' sheet into friction-page-survey.xlsx: method, live-formula summary,
the essay's figures, the EI 2.0 coding dictionary and every coded answer."""
import re
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from scipy.stats import chi2_contingency

XLSX, ROWS, DICT = sys.argv[1], sys.argv[2], sys.argv[3]
rows = pd.read_pickle(ROWS)
d = pd.read_pickle(DICT)
WORDS, CORES = d['WORDS'], d['CORE_ORDER']

wb = load_workbook(XLSX)
if 'anger-at-work' in wb.sheetnames:
    del wb['anger-at-work']
ws = wb.create_sheet('anger-at-work')

NAVY = PatternFill('solid', fgColor='002D61')
SOFT = PatternFill('solid', fgColor='EAF1F8')
WHITE_B = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)
H1 = Font(bold=True, size=15, color='002D61')
H2 = Font(bold=True, size=12, color='002D61')
NOTE = Font(italic=True, color='5A6A82')
WRAP = Alignment(wrap_text=True, vertical='top')
thin = Border(bottom=Side(style='thin', color='DCE6F1'))

def header(r, labels, c0=1):
    for i, lab in enumerate(labels):
        cell = ws.cell(r, c0 + i, lab)
        cell.fill, cell.font, cell.alignment = NAVY, WHITE_B, Alignment(wrap_text=True, vertical='center')

# ---- layout: where the coded data table lands ----
n = len(rows)
r = 1
ws.cell(r, 1, 'Anger at work: which emotion feels most off-limits? (Q6)').font = H1
ws.cell(r + 1, 1, 'Source: friction-page survey (The American Emotions Survey), Sheet1, Q6. 1,188 working U.S. adults, Oct 2025 – Jan 2026. '
                  'Summary figures below are live formulas over the coded answers at the bottom of this sheet.').font = NOTE
r = 4
ws.cell(r, 1, 'How each answer was coded').font = H2
steps = [
    '1. Q6 is open-ended: "Which emotion feels most off-limits in your workplace (e.g., anger, sadness, joy, vulnerability) — and why?"',
    '2. Only the emotion the person NAMED is coded: the text before "because", "since", "when", "if" or the first full stop. '
    'Example: "Mostly joy because others are angry and bitter" is coded joy, not anger.',
    '3. The named emotion is matched to the five core emotions in TalentSmartEQ\'s Emotional Intelligence 2.0 feeling-words table '
    '(Happy, Sad, Angry, Afraid, Ashamed), with high / medium / low intensity. Rage, frustration and disgust/disdain are forms of anger. '
    'Vulnerability is not in the table; it is grouped with Afraid (insecure, exposed).',
    '4. Answers that name no emotion ("none", "I work alone", "all emotions") are not counted. Percentages are of people who named an emotion. '
    'Someone can name more than one, so the five cores add to more than 100%.',
    '5. Role and gender gaps were tested with a chi-square test (p < .05 = statistically significant).',
]
for i, s in enumerate(steps):
    c = ws.cell(r + 1 + i, 1, s); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r + 1 + i, start_column=1, end_row=r + 1 + i, end_column=14)
    ws.row_dimensions[r + 1 + i].height = 30

SUM_R = r + len(steps) + 2           # summary title row
TOT_R = SUM_R + 2 + len(CORES)       # totals under the summary
E_R = TOT_R + 4                      # essay figures title row (8 figures follow)
DICT_R = E_R + 2 + 8 + 2             # dictionary title row
DATA_H = DICT_R + len(WORDS) + 4     # data header row
D1, D2 = DATA_H + 1, DATA_H + n      # first/last data row
col = {name: get_column_letter(i + 1) for i, name in enumerate(rows.columns)}
rng = lambda name: f"${col[name]}${D1}:${col[name]}${D2}"

# ---- summary table (formulas) ----
ws.cell(SUM_R, 1, 'Summary by core emotion').font = H2
labels = ['Core emotion', 'Named it (n)', '% of those who named an emotion', '% of all respondents',
          'Non-leaders', 'Leaders', 'Role gap p-value', 'Women', 'Men', 'Gender gap p-value']
header(SUM_R + 1, labels)
named_yes = rows['Named an emotion'] == 'Yes'
def pval(mask, a, b):
    tab = [[(mask & named_yes & a).sum(), (~mask & named_yes & a).sum()], [(mask & named_yes & b).sum(), (~mask & named_yes & b).sum()]]
    return float(chi2_contingency(tab)[1])
for i, c in enumerate(CORES):
    rr = SUM_R + 2 + i
    m = rows[c] == 1
    cr = rng(c)
    ws.cell(rr, 1, c).font = BOLD
    ws.cell(rr, 2, f'=SUM({cr})')
    ws.cell(rr, 3, f'=SUM({cr})/COUNTIF({rng("Named an emotion")},"Yes")')
    ws.cell(rr, 4, f'=SUM({cr})/COUNTA({rng("ResponseId")})')
    for j, (grpcol, val) in enumerate([('Leader role', 'Non-leaders'), ('Leader role', 'Leaders')]):
        ws.cell(rr, 5 + j, f'=COUNTIFS({cr},1,{rng(grpcol)},"{val}")/COUNTIFS({rng("Named an emotion")},"Yes",{rng(grpcol)},"{val}")')
    ws.cell(rr, 7, pval(m, rows['Leader role'] == 'Non-leaders', rows['Leader role'] == 'Leaders'))
    for j, val in enumerate(['Female', 'Male']):
        ws.cell(rr, 8 + j, f'=COUNTIFS({cr},1,{rng("Gender")},"{val}")/COUNTIFS({rng("Named an emotion")},"Yes",{rng("Gender")},"{val}")')
    ws.cell(rr, 10, pval(m, rows['Gender'] == 'Female', rows['Gender'] == 'Male'))
    for cc in (3, 4, 5, 6, 8, 9):
        ws.cell(rr, cc).number_format = '0.0%'
    for cc in (7, 10):
        ws.cell(rr, cc).number_format = '0.0000'
    for cc in range(1, 11):
        ws.cell(rr, cc).border = thin
tot = TOT_R
ws.cell(tot, 1, 'People who named an emotion').font = BOLD
ws.cell(tot, 2, f'=COUNTIF({rng("Named an emotion")},"Yes")')
ws.cell(tot + 1, 1, 'All respondents').font = BOLD
ws.cell(tot + 1, 2, f'=COUNTA({rng("ResponseId")})')
ws.cell(tot + 2, 1, 'p-values are static results of a chi-square test run in Python on the same coded rows.').font = NOTE

# ---- figures used in the essay ----
E = E_R
ws.cell(E, 1, 'Figures used in the "Anger is off-limits at work" essay').font = H2
header(E + 1, ['Figure', 'Value', 'Essay wording'])
vf = f'(ISNUMBER(SEARCH(", Vulnerability,",", "&{rng("Specific emotions")}&","))+ISNUMBER(SEARCH(", Fear,",", "&{rng("Specific emotions")}&","))>0)'
def vf_share(g):
    return (f'SUMPRODUCT(({rng("Named an emotion")}="Yes")*({rng("Gender")}="{g}")*{vf})'
            f'/COUNTIFS({rng("Named an emotion")},"Yes",{rng("Gender")},"{g}")')
ang = SUM_R + 2  # Angry is the first core row
figs = [
    ('Anger, of those who named an emotion', f'=C{ang}', '67%', '0.0%'),
    ('Anger, non-leaders', f'=E{ang}', '74%', '0.0%'),
    ('Anger, leaders', f'=F{ang}', '59%', '0.0%'),
    ('Anger, women', f'=H{ang}', '72%', '0.0%'),
    ('Anger, men', f'=I{ang}', '62%', '0.0%'),
    ('Vulnerability or fear, women', '=' + vf_share('Female'), '', '0.0%'),
    ('Vulnerability or fear, men', '=' + vf_share('Male'), '', '0.0%'),
    ('Men vs. women, vulnerability or fear', f'=B{E + 8}/B{E + 7}', '"nearly twice as likely"', '0.00"x"'),
]
assert E + 2 + len(figs) < DICT_R, 'essay figures would overlap the dictionary'
for i, (lab, f, words, fmt) in enumerate(figs):
    rr = E + 2 + i
    ws.cell(rr, 1, lab); ws.cell(rr, 2, f).number_format = fmt; ws.cell(rr, 3, words)
    for cc in range(1, 4):
        ws.cell(rr, cc).border = thin

# ---- coding dictionary ----
ws.cell(DICT_R, 1, 'Coding dictionary (Emotional Intelligence 2.0 core emotions)').font = H2
header(DICT_R + 1, ['Core emotion', 'Intensity', 'Specific emotion', 'Words that count (word starts)', 'Answers coded'])
def stems(rx):
    s = re.sub(r'\\b|\(|\)|\\s|\?', '', rx).replace('.', ' ')
    return ', '.join(x for x in s.split('|') if x)
for i, (w, c, inten, rx) in enumerate(WORDS):
    rr = DICT_R + 2 + i
    ws.cell(rr, 1, c); ws.cell(rr, 2, inten); ws.cell(rr, 3, w); ws.cell(rr, 4, stems(rx))
    ws.cell(rr, 5, f'=SUMPRODUCT(--ISNUMBER(SEARCH(", {w},",", "&{rng("Specific emotions")}&",")))')
    for cc in range(1, 6):
        ws.cell(rr, cc).border = thin
        ws.cell(rr, cc).alignment = Alignment(wrap_text=True, vertical='top')

# ---- coded answers ----
ws.cell(DATA_H - 1, 1, f'Every answer, as coded ({n:,} rows). Filter or sort with the arrows in the header.').font = H2
header(DATA_H, list(rows.columns))
for i, rec in enumerate(rows.itertuples(index=False)):
    for j, v in enumerate(rec):
        cell = ws.cell(D1 + i, j + 1, None if (isinstance(v, float) and pd.isna(v)) else v)
        if j in (3, 4):
            cell.alignment = WRAP
ws.auto_filter.ref = f'A{DATA_H}:{get_column_letter(len(rows.columns))}{D2}'

widths = {'A': 30, 'B': 14, 'C': 12, 'D': 60, 'E': 34, 'F': 12, 'G': 9, 'H': 9, 'I': 9, 'J': 9, 'K': 9, 'L': 26, 'M': 16, 'N': 18}
for k, v in widths.items():
    ws.column_dimensions[k].width = v
for rr in (SUM_R + 1, DICT_R + 1, E + 1, DATA_H):
    ws.row_dimensions[rr].height = 32
ws.sheet_view.showGridLines = False
ws.freeze_panes = None
wb.save(XLSX)
print('saved; summary rows', SUM_R, 'essay figs', E, 'dictionary', DICT_R, 'data', D1, '-', D2)
